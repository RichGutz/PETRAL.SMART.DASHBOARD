import os
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- STYLING HELPERS ---

NAVY_HEX = "0F2C59"
BLUE_HEX = "1E56A0"
LIGHT_BLUE_HEX = "EBF5FB"
GOLD_HEX = "D4AF37"
GOLD_LIGHT_HEX = "FEF9E7"
GRAY_BORDER_HEX = "CBD5E1"
GRAY_BG_HEX = "F8FAFC"
TEXT_DARK_HEX = "1E293B"
TEXT_MUTED_HEX = "64748B"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_padding(cell, top=120, bottom=120, left=160, right=160):
    # values in dxa (1 pt = 20 dxa)
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_callout_border(cell, bg_color=GOLD_LIGHT_HEX, border_color=GOLD_HEX):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_table_borders(table, border_color=GRAY_BORDER_HEX, sz="4"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_header_run(paragraph, text, size=11, bold=True, color_rgb=(15, 44, 89)):
    run = paragraph.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    return run

def add_body_run(paragraph, text, size=9.5, bold=False, italic=False, color_rgb=(30, 41, 59)):
    run = paragraph.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    return run

# --- EXTRACTION OF HTML SLIDES ---

def parse_presentation_html(html_path):
    with open(html_path, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    slides_data = []
    slides = soup.find_all(class_="slide-page")
    for idx, s in enumerate(slides):
        slide = {
            "slide_num": idx + 1,
            "tagline": "",
            "title": "",
            "subtitle": "",
            "headers": [],
            "rows": [],
            "callout_title": "",
            "callout_body": "",
            "footer": ""
        }
        
        # Header
        header = s.find(class_="slide-header")
        if header:
            t_el = header.find(class_="tagline")
            if t_el:
                slide["tagline"] = t_el.get_text(strip=True)
            h2_el = header.find("h2")
            if h2_el:
                slide["title"] = h2_el.get_text(strip=True)
            sub_el = header.find(class_="subtitle")
            if sub_el:
                slide["subtitle"] = sub_el.get_text(strip=True)
                
        # Body
        body = s.find(class_="slide-body")
        if body:
            table = body.find("table")
            if table:
                # Headers
                th_list = table.find_all("th")
                slide["headers"] = [th.get_text(" ", strip=True) for th in th_list]
                
                # Rows
                tbody = table.find("tbody") or table
                for tr in tbody.find_all("tr"):
                    tds = tr.find_all("td")
                    if tds:
                        row_items = []
                        for td in tds:
                            # Let's inspect contents: strong, br, spans, text
                            td_soup = BeautifulSoup(str(td), "html.parser").find("td")
                            
                            # Clean bullet lines
                            for br in td_soup.find_all("br"):
                                br.replace_with("\n")
                            raw_lines = [l.strip() for l in td_soup.get_text().split("\n") if l.strip()]
                            
                            # Also check if there's a strong element
                            strong_tag = td.find("strong")
                            strong_txt = strong_tag.get_text(strip=True) if strong_tag else ""
                            
                            badge = td.find(class_=lambda c: c and "badge" in c)
                            badge_txt = badge.get_text(strip=True) if badge else ""
                            
                            row_items.append({
                                "text": "\n".join(raw_lines),
                                "lines": raw_lines,
                                "strong": strong_txt,
                                "badge": badge_txt
                            })
                        slide["rows"].append(row_items)
                        
            # Callouts
            callout = body.find(class_="card-callout")
            if callout:
                callout_copy = BeautifulSoup(str(callout), "html.parser").find(class_="card-callout")
                strong_c = callout_copy.find("strong")
                if strong_c:
                    slide["callout_title"] = strong_c.get_text(strip=True)
                    strong_c.decompose()
                slide["callout_body"] = callout_copy.get_text(" ", strip=True)
                
            # Footer
            for div in body.find_all("div", recursive=False):
                if "NAVIERA PETRAL" in div.get_text():
                    slide["footer"] = div.get_text(strip=True)

        slides_data.append(slide)
    return slides_data

# --- BUILD WORD DOCUMENT ---

def build_word_document(slides_data, output_docx_path, logo_path=None):
    doc = Document()
    
    # 1. Page Setup: Landscape Orientation for executive wide tables
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        
    # --- COVER / INSTRUCTIONS SECTION ---
    
    # Header logo or title
    if logo_path and os.path.exists(logo_path):
        try:
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_logo.paragraph_format.space_after = Pt(4)
            run_logo = p_logo.add_run()
            run_logo.add_picture(logo_path, width=Inches(2.2))
        except Exception as e:
            print(f"Note: Could not add logo: {e}")

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    add_header_run(p_title, "TRILOGÍA PETRAL: ONE, TWO & THREE", size=22, bold=True, color_rgb=(15, 44, 89))
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(10)
    add_body_run(p_sub, "GUION EDITABLE DE TEXTOS PARA EL PROJECT MANAGER (PM)", size=12.5, bold=True, color_rgb=(30, 86, 160))
    add_body_run(p_sub, "  |  Presentación Ejecutiva & Capacitación Oficial", size=10.5, italic=True, color_rgb=(100, 116, 139))
    
    # Instructions Box for PM
    t_inst = doc.add_table(rows=1, cols=1)
    t_inst.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_inst = t_inst.cell(0, 0)
    set_cell_background(c_inst, "F0F7FF")
    set_callout_border(c_inst, bg_color="F0F7FF", border_color="1E56A0")
    set_cell_padding(c_inst, top=140, bottom=140, left=200, right=200)
    
    p_inst_h = c_inst.paragraphs[0]
    p_inst_h.paragraph_format.space_after = Pt(4)
    add_header_run(p_inst_h, "✍️ INSTRUCCIONES PARA EL PROJECT MANAGER (PM) - EDICIÓN DIRECTA", size=11, bold=True, color_rgb=(30, 86, 160))
    
    p_inst_b = c_inst.add_paragraph()
    p_inst_b.paragraph_format.space_after = Pt(0)
    inst_text = (
        "1. Edición de Textos: Puede modificar directamente en este archivo Word cualquier texto de las 6 diapositivas (títulos, subtítulos, contenidos de las celdas, viñetas explicativas y mensajes clave).\n"
        "2. Respeto de Estructura: No elimine las filas ni altere los marcadores [DIAPOSITIVA X] o [TABLA PRINCIPAL], ya que nuestro sistema automático lee este documento para sincronizarlo con el visor web en producción (petral.one.geeksoft.tech).\n"
        "3. Finalización: Una vez editados y guardados los cambios en este archivo, simplemente devuélvalo al equipo técnico para su re-inyección automática e instantánea."
    )
    add_body_run(p_inst_b, inst_text, size=9.5, color_rgb=(30, 41, 59))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # --- SLIDES (1 TO 6) ---
    for s_idx, slide in enumerate(slides_data):
        if s_idx > 0:
            doc.add_page_break()
            
        # Slide Title Section
        p_slide_title = doc.add_paragraph()
        p_slide_title.paragraph_format.space_before = Pt(4)
        p_slide_title.paragraph_format.space_after = Pt(2)
        add_header_run(p_slide_title, f"[DIAPOSITIVA {slide['slide_num']}] — {slide['title']}", size=15, bold=True, color_rgb=(15, 44, 89))
        
        # Metadata / Tagline / Subtitle Box
        t_meta = doc.add_table(rows=2, cols=2)
        t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_meta, border_color="E2E8F0", sz="4")
        
        # Col widths: 20% and 80%
        for row in t_meta.rows:
            row.cells[0].width = Inches(1.8)
            row.cells[1].width = Inches(8.0)
            
        # Row 1: Tagline
        c00 = t_meta.cell(0, 0)
        c01 = t_meta.cell(0, 1)
        set_cell_background(c00, "F8FAFC")
        set_cell_padding(c00, top=80, bottom=80, left=120, right=120)
        set_cell_padding(c01, top=80, bottom=80, left=120, right=120)
        p00 = c00.paragraphs[0]
        p00.paragraph_format.space_after = Pt(0)
        add_header_run(p00, "Tagline / Badge:", size=9.5, bold=True, color_rgb=(30, 86, 160))
        p01 = c01.paragraphs[0]
        p01.paragraph_format.space_after = Pt(0)
        add_body_run(p01, slide["tagline"], size=9.5, bold=True, color_rgb=(30, 86, 160))
        
        # Row 2: Subtitle
        c10 = t_meta.cell(1, 0)
        c11 = t_meta.cell(1, 1)
        set_cell_background(c10, "F8FAFC")
        set_cell_padding(c10, top=80, bottom=80, left=120, right=120)
        set_cell_padding(c11, top=80, bottom=80, left=120, right=120)
        p10 = c10.paragraphs[0]
        p10.paragraph_format.space_after = Pt(0)
        add_header_run(p10, "Subtítulo / Bajada:", size=9.5, bold=True, color_rgb=(100, 116, 139))
        p11 = c11.paragraphs[0]
        p11.paragraph_format.space_after = Pt(0)
        add_body_run(p11, slide["subtitle"], size=9.5, italic=True, color_rgb=(71, 85, 105))
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # Table Section Header
        p_tbl_lbl = doc.add_paragraph()
        p_tbl_lbl.paragraph_format.space_before = Pt(4)
        p_tbl_lbl.paragraph_format.space_after = Pt(2)
        add_header_run(p_tbl_lbl, f"[TABLA PRINCIPAL — DIAPOSITIVA {slide['slide_num']}]", size=10.5, bold=True, color_rgb=(15, 44, 89))
        
        # Main Content Table
        headers = slide["headers"]
        rows = slide["rows"]
        
        num_cols = len(headers) if headers else (len(rows[0]) if rows else 1)
        num_rows = len(rows) + (1 if headers else 0)
        
        t_main = doc.add_table(rows=num_rows, cols=num_cols)
        t_main.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_main, border_color=GRAY_BORDER_HEX, sz="4")
        
        # Distribute column widths across available width (~9.8 inches)
        col_width = Inches(9.8 / num_cols)
        for row in t_main.rows:
            for cell in row.cells:
                cell.width = col_width
                
        # Header Row
        current_r = 0
        if headers:
            for col_idx, h_text in enumerate(headers):
                cell = t_main.cell(current_r, col_idx)
                set_cell_background(cell, NAVY_HEX)
                set_cell_padding(cell, top=100, bottom=100, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == num_cols - 1 and num_cols > 3 else WD_ALIGN_PARAGRAPH.LEFT
                add_header_run(p, h_text, size=9.5, bold=True, color_rgb=(255, 255, 255))
            current_r += 1
            
        # Data Rows
        for r_data in rows:
            for col_idx, item in enumerate(r_data):
                if col_idx < num_cols:
                    cell = t_main.cell(current_r, col_idx)
                    set_cell_padding(cell, top=90, bottom=90, left=120, right=120)
                    
                    # Alternating row background for subtle elegance
                    if current_r % 2 == 0:
                        set_cell_background(cell, GRAY_BG_HEX)
                        
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    
                    lines = item["lines"]
                    if lines:
                        for l_idx, line in enumerate(lines):
                            if l_idx > 0:
                                p = cell.add_paragraph()
                                p.paragraph_format.space_after = Pt(2)
                                
                            # Check if line matches strong title or badge
                            if l_idx == 0 and item["strong"] and item["strong"] in line:
                                add_body_run(p, line, size=9, bold=True, color_rgb=(15, 44, 89))
                            elif item["badge"] and item["badge"] in line:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                add_body_run(p, f"[{line}]", size=9, bold=True, color_rgb=(30, 86, 160))
                            elif line.startswith("•") or line.startswith("-"):
                                add_body_run(p, line, size=8.5, color_rgb=(30, 41, 59))
                            else:
                                add_body_run(p, line, size=9, color_rgb=(30, 41, 59))
                    else:
                        add_body_run(p, "", size=9)
            current_r += 1
            
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # Callout Card / Mensaje Clave
        if slide["callout_body"]:
            t_call = doc.add_table(rows=1, cols=1)
            t_call.alignment = WD_TABLE_ALIGNMENT.CENTER
            c_call = t_call.cell(0, 0)
            set_callout_border(c_call, bg_color=GOLD_LIGHT_HEX, border_color=GOLD_HEX)
            set_cell_padding(c_call, top=100, bottom=100, left=160, right=160)
            
            p_call = c_call.paragraphs[0]
            p_call.paragraph_format.space_after = Pt(0)
            
            call_title = slide["callout_title"] if slide["callout_title"] else "Mensaje Clave:"
            add_header_run(p_call, f"💡 {call_title} ", size=9.5, bold=True, color_rgb=(15, 44, 89))
            add_body_run(p_call, slide["callout_body"], size=9, color_rgb=(30, 41, 59))
            
        # Footer Note (if present)
        if slide.get("footer"):
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            p_ft = doc.add_paragraph()
            p_ft.paragraph_format.space_before = Pt(2)
            p_ft.paragraph_format.space_after = Pt(0)
            p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_body_run(p_ft, slide["footer"], size=8, italic=True, color_rgb=(148, 163, 184))

    # Save document
    doc.save(output_docx_path)
    print(f"Word document generated successfully at: {output_docx_path}")

if __name__ == "__main__":
    html_file = "Proposal/PRESENTACION_TRILOGIA_PETRAL_ONE_TWO_THREE.html"
    output_docx = "Proposal/GUION_TEXTOS_PRESENTACION_SOTA_PETRAL.docx"
    logo_file = "Proposal/LOGO.PETRAL.ONE.HORIZONTAL.png"
    
    slides = parse_presentation_html(html_file)
    print(f"Extracted {len(slides)} slides.")
    build_word_document(slides, output_docx, logo_path=logo_file)
